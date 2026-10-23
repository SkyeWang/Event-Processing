# crawler/spider.py

import re
import time
import random
import string
import openpyxl as op
import selenium.common.exceptions
import requests
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options as ChromeOptions
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from datetime import datetime, timedelta
from bs4 import BeautifulSoup

# 随机 User-Agent 列表
USER_AGENTS = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/112.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:109.0) Gecko/20100101 Firefox/109.0",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/15.1 Safari/605.1.15"
]

# 代理列表 - Twitter爬虫专用
PROXIES = []  # 填写 proxy 时使用，或使用 get_twitter_proxies() 动态获取

def get_twitter_proxies(protocol='http', count=10, country_code='US'):
    """
    从 https://proxy.scdn.io/ 获取适合Twitter的免费代理
    
    Args:
        protocol: 代理协议 (http, https, socks4, socks5)
        count: 获取数量 (最大20)
        country_code: 国家代码 (US, CN, UK等)
    
    Returns:
        list: 格式化的代理列表
    """
    try:
        url = 'https://proxy.scdn.io/api/get_proxy.php'
        params = {
            'protocol': protocol,
            'count': min(count, 20),  # API限制最大20个
            'country_code': country_code
        }
        response = requests.get(url, params=params, timeout=10)
        data = response.json()
        
        if data.get('code') == 200:
            proxies = data.get('data', {}).get('proxies', [])
            # 添加协议前缀
            formatted_proxies = []
            for proxy in proxies:
                if protocol == 'http':
                    formatted_proxies.append(f"http://{proxy}")
                elif protocol == 'https':
                    formatted_proxies.append(f"https://{proxy}")
                elif protocol == 'socks4':
                    formatted_proxies.append(f"socks4://{proxy}")
                elif protocol == 'socks5':
                    formatted_proxies.append(f"socks5://{proxy}")
                else:
                    formatted_proxies.append(f"http://{proxy}")
            return formatted_proxies
        else:
            print(f"获取Twitter代理失败: {data.get('message', '未知错误')}")
            return []
    except Exception as e:
        print(f"获取Twitter代理时发生错误: {e}")
        return []

def update_twitter_proxies():
    """更新Twitter爬虫代理列表"""
    global PROXIES
    print("正在获取Twitter专用代理...")
    # 优先获取美国HTTP代理，Twitter访问更稳定
    PROXIES = get_twitter_proxies('http', 10, 'US')
    if not PROXIES:
        # 如果没有美国代理，获取任意国家的HTTP代理
        print("未获取到美国代理，尝试获取任意国家代理...")
        PROXIES = get_twitter_proxies('http', 10, 'all')
    if not PROXIES:
        # 最后尝试SOCKS5代理
        print("HTTP代理获取失败，尝试SOCKS5代理...")
        PROXIES = get_twitter_proxies('socks5', 10, 'all')
    
    print(f"获取到 {len(PROXIES)} 个Twitter代理")
    return PROXIES

def get_random_proxy():
    return random.choice(PROXIES) if PROXIES else None

# 全局技术关键词，由 run_crawler 或 run_crawler_stream 覆盖
TECH_KEYWORDS = []

# Excel 文件路径
file_path = '8.xlsx'

def is_probably_english(text, threshold=0.7):
    total = sum(c.isalpha() for c in text)
    eng = sum(c in string.ascii_letters for c in text)
    return False if total == 0 else (eng / total) >= threshold


def is_english_and_relevant(text):
    if not is_probably_english(text):
        return False, None
    norm = re.sub(r'[^a-z0-9\s]', ' ', text.lower())
    norm = ' '.join(norm.split())
    matched = [kw for kw in TECH_KEYWORDS if kw in norm]
    return (True, ', '.join(matched)) if matched else (False, None)

class TwitterScraper:
    def __init__(self, keywords, proxy=None):
        options = ChromeOptions()
        options.add_argument(f"user-agent={random.choice(USER_AGENTS)}")
        if proxy:
            options.add_argument(f"--proxy-server={proxy}")
        options.add_argument("--disable-blink-features=AutomationControlled")
        self.driver = webdriver.Chrome(options=options)
        self.wait = WebDriverWait(self.driver, 30)
        self.keywords = [kw.lower() for kw in keywords]
        self.workbook = op.load_workbook(file_path)
        self.sheet = self.workbook.active
        self.count = self.sheet.max_row + 1

    def simulate_mouse_movements(self):
        actions = ActionChains(self.driver)
        for _ in range(random.randint(5, 10)):
            actions.move_by_offset(random.randint(-50, 50), random.randint(-50, 50)).pause(random.uniform(0.2, 0.5))
        try:
            actions.perform()
        except Exception:
            pass

    def check_retry(self):
        if 'retry' in self.driver.page_source.lower():
            time.sleep(30)
            self.driver.refresh()
            time.sleep(random.uniform(5, 8))

    def login(self, username, password):
        self.driver.get('https://twitter.com/login')
        time.sleep(random.uniform(3, 5))
        user_field = self.wait.until(EC.presence_of_element_located((By.NAME, 'text')))
        user_field.send_keys(username, Keys.RETURN)
        time.sleep(random.uniform(4, 6))
        pwd_field = self.wait.until(EC.presence_of_element_located((By.NAME, 'password')))
        pwd_field.send_keys(password, Keys.RETURN)
        time.sleep(random.uniform(4, 6))
        self.simulate_mouse_movements()

    def scrape_tweets(self, start_date, end_date=None):
        end_date = end_date or datetime.today()
        current = start_date
        while current <= end_date:
            next_day = current + timedelta(days=1)
            self._search_and_scroll(current, next_day)
            current = next_day
            self.workbook.save(file_path)

    def _search_and_scroll(self, since, until):
        sin = since.strftime('%Y-%m-%d')
        unt = until.strftime('%Y-%m-%d')
        query = '+'.join(self.keywords)
        url = f'https://twitter.com/search?q={query}+since%3A{sin}+until%3A{unt}&src=typed_query'
        self.driver.get(url)
        time.sleep(random.uniform(2, 4))
        self.check_retry()

        last_pos = -1
        attempts = 0
        while attempts < 3:
            self.simulate_mouse_movements()
            self.check_retry()
            soup = BeautifulSoup(self.driver.page_source, 'html.parser')
            for art in soup.find_all('article'):
                try:
                    txt, matched_info = art.find('div', {'data-testid': 'tweetText'}).text, None
                    valid, matched = is_english_and_relevant(txt)
                    if not valid:
                        continue
                    self.sheet.cell(row=self.count, column=1, value=txt)
                    self.sheet.cell(row=self.count, column=2, value=matched)
                    self.count += 1
                except Exception:
                    continue

            self.driver.execute_script('window.scrollTo(0, document.body.scrollHeight);')
            time.sleep(random.uniform(2, 4))
            self.check_retry()

            pos = self.driver.execute_script('return window.pageYOffset;')
            if pos == last_pos:
                attempts += 1
            else:
                last_pos = pos
                attempts = 0

    # 实时流式爬取
    def stream_tweets(self, start_date, end_date=None):
        end_date = end_date or datetime.today()
        current = start_date
        while current <= end_date:
            next_day = current + timedelta(days=1)
            for txt, matched in self._stream_search_and_scroll(current, next_day):
                yield txt, matched
            current = next_day
            self.workbook.save(file_path)

    def _stream_search_and_scroll(self, since, until):
        sin = since.strftime('%Y-%m-%d')
        unt = until.strftime('%Y-%m-%d')
        query = '+'.join(self.keywords)
        url = f'https://twitter.com/search?q={query}+since%3A{sin}+until%3A{unt}&src=typed_query'
        self.driver.get(url)
        time.sleep(random.uniform(2, 4))
        self.check_retry()

        last_pos = -1
        attempts = 0
        while attempts < 3:
            self.simulate_mouse_movements()
            self.check_retry()
            soup = BeautifulSoup(self.driver.page_source, 'html.parser')
            for art in soup.find_all('article'):
                try:
                    txt = art.find('div', {'data-testid': 'tweetText'}).text
                    valid, matched = is_english_and_relevant(txt)
                    if valid:
                        self.sheet.cell(row=self.count, column=1, value=txt)
                        self.sheet.cell(row=self.count, column=2, value=matched)
                        self.count += 1
                        yield txt, matched
                except Exception:
                    continue

            self.driver.execute_script('window.scrollTo(0, document.body.scrollHeight);')
            time.sleep(random.uniform(2, 4))
            self.check_retry()

            pos = self.driver.execute_script('return window.pageYOffset;')
            if pos == last_pos:
                attempts += 1
            else:
                last_pos, attempts = pos, 0


def run_crawler(keywords, start_date, end_date, username, password, proxy=None, auto_proxy=False):
    global TECH_KEYWORDS
    TECH_KEYWORDS = [kw.lower() for kw in keywords]
    
    # 如果启用自动代理且没有提供代理，则自动获取
    if auto_proxy and not proxy:
        update_twitter_proxies()
        proxy = get_random_proxy()
    
    scraper = TwitterScraper(keywords, proxy)
    scraper.login(username, password)
    scraper.scrape_tweets(start_date, end_date)
    scraper.driver.quit()
    wb = op.load_workbook(file_path)
    sheet = wb.active
    return [
        {'text': row[0], 'matched': row[1]}
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if row[0]
    ]


def run_crawler_stream(keywords, start_date, end_date, username, password, proxy=None, auto_proxy=False):
    global TECH_KEYWORDS
    TECH_KEYWORDS = [kw.lower() for kw in keywords]
    
    # 如果启用自动代理且没有提供代理，则自动获取
    if auto_proxy and not proxy:
        update_twitter_proxies()
        proxy = get_random_proxy()
    
    scraper = TwitterScraper(keywords, proxy)
    scraper.login(username, password)
    for txt, matched in scraper.stream_tweets(start_date, end_date):
        yield {'text': txt, 'matched': matched}
    scraper.driver.quit()

